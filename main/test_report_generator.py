from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase

from main.models import (
    ProgramOutcome, Course, LearningOutcome,
    Assessment, AssessmentToLO, LOToPO, Student, StudentGrade
)
from main.services.report_generator import CourseReportService


class CourseReportTestCase(TestCase):
    def setUp(self):
        # Same setup as the slide's worked example (see POCalculatorTestCase)
        self.po1 = ProgramOutcome.objects.create(code="PO1", title="PO 1", description="Desc 1")
        self.po2 = ProgramOutcome.objects.create(code="PO2", title="PO 2", description="Desc 2")
        # Not mapped to the course, must not appear in the report
        ProgramOutcome.objects.create(code="PO3", title="PO 3", description="Desc 3")

        self.course = Course.objects.create(code="CSE 311", title="Software")
        self.lo1 = LearningOutcome.objects.create(course=self.course, code="LO1", description="LO 1")
        self.lo2 = LearningOutcome.objects.create(course=self.course, code="LO2", description="LO 2")

        self.midterm = Assessment.objects.create(course=self.course, name="Midterm", weight=60.0)
        self.project = Assessment.objects.create(course=self.course, name="Project", weight=40.0)

        AssessmentToLO.objects.create(assessment=self.midterm, learning_outcome=self.lo1, weight=60.0)
        AssessmentToLO.objects.create(assessment=self.midterm, learning_outcome=self.lo2, weight=40.0)
        AssessmentToLO.objects.create(assessment=self.project, learning_outcome=self.lo1, weight=40.0)
        AssessmentToLO.objects.create(assessment=self.project, learning_outcome=self.lo2, weight=100.0)

        LOToPO.objects.create(learning_outcome=self.lo1, program_outcome=self.po1, weight=100.0)
        LOToPO.objects.create(learning_outcome=self.lo2, program_outcome=self.po1, weight=50.0)
        LOToPO.objects.create(learning_outcome=self.lo2, program_outcome=self.po2, weight=100.0)

        # Elif: strong student (PO1 = 85.0, PO2 = 82.86)
        self.elif_ = Student.objects.create(student_number="1001", first_name="Elif", last_name="Cato")
        StudentGrade.objects.create(student=self.elif_, assessment=self.midterm, score=90.0)
        StudentGrade.objects.create(student=self.elif_, assessment=self.project, score=80.0)

        # Can: weak student (PO1 = 45.05, PO2 = 47.14)
        self.can = Student.objects.create(student_number="1002", first_name="Can", last_name="Yilmaz")
        StudentGrade.objects.create(student=self.can, assessment=self.midterm, score=40.0)
        StudentGrade.objects.create(student=self.can, assessment=self.project, score=50.0)

        # No grades in this course, must not appear in the report
        Student.objects.create(student_number="1003", first_name="Ada", last_name="Kaya")

    def test_report_contains_only_graded_students_and_mapped_pos(self):
        report = CourseReportService.build_course_report(self.course)

        self.assertEqual([row["student"] for row in report["students"]], [self.elif_, self.can])
        self.assertEqual(report["po_codes"], ["PO1", "PO2"])
        self.assertEqual(report["lo_codes"], ["LO1", "LO2"])
        self.assertEqual(report["students"][0]["grades"], [90.0, 80.0])

    def test_po_summary_statistics(self):
        report = CourseReportService.build_course_report(self.course, threshold=60.0)
        po1 = report["po_summary"][0]

        self.assertEqual(po1["code"], "PO1")
        self.assertEqual(po1["count"], 2)
        self.assertAlmostEqual(po1["average"], (85.0 + 45.05) / 2, places=1)
        self.assertAlmostEqual(po1["min"], 45.05, places=1)
        self.assertAlmostEqual(po1["max"], 85.0, places=1)
        self.assertEqual(po1["achieved"], 1)
        self.assertEqual(po1["achievement_rate"], 50.0)

    def test_at_risk_students_depend_on_threshold(self):
        report = CourseReportService.build_course_report(self.course, threshold=60.0)
        self.assertEqual(len(report["at_risk"]), 1)
        self.assertEqual(report["at_risk"][0]["student"], self.can)
        self.assertEqual(report["at_risk"][0]["weak_pos"], ["PO1", "PO2"])

        # With a stricter threshold Elif's PO2 (82.86) is also below
        report = CourseReportService.build_course_report(self.course, threshold=84.0)
        self.assertEqual([item["student"] for item in report["at_risk"]], [self.can, self.elif_])
        self.assertEqual(report["at_risk"][1]["weak_pos"], ["PO2"])

    def test_course_without_grades_produces_empty_report(self):
        empty_course = Course.objects.create(code="MAT 241", title="Linear Algebra")
        report = CourseReportService.build_course_report(empty_course)

        self.assertEqual(report["students"], [])
        self.assertEqual(report["at_risk"], [])
        # The workbook can still be generated without errors
        workbook = openpyxl.load_workbook(CourseReportService.export_excel(empty_course))
        self.assertEqual(workbook["At-Risk Students"]["A2"].value, "No students are below the threshold.")

    def test_excel_export_sheets_and_values(self):
        workbook = openpyxl.load_workbook(CourseReportService.export_excel(self.course))
        self.assertEqual(workbook.sheetnames, ["Summary", "Students", "At-Risk Students"])

        students = workbook["Students"]
        headers = [cell.value for cell in students[1]]
        self.assertEqual(headers, ["Student No", "Name", "Midterm", "Project", "LO1", "LO2", "PO1", "PO2"])
        self.assertEqual(students["A2"].value, "1001")
        self.assertAlmostEqual(students["G2"].value, 85.0, places=1)
        self.assertEqual(students["A4"].value, "Class Average")

        at_risk = workbook["At-Risk Students"]
        self.assertEqual(at_risk["A2"].value, "1002")
        self.assertEqual(at_risk["C2"].value, "PO1, PO2")

    def test_management_command_writes_file(self):
        with TemporaryDirectory() as tmp:
            output = Path(tmp) / "report.xlsx"
            call_command("export_course_report", "CSE 311", "--output", str(output), stdout=StringIO())
            self.assertTrue(output.exists())
            self.assertIn("Summary", openpyxl.load_workbook(output).sheetnames)

    def test_management_command_unknown_course(self):
        with self.assertRaises(CommandError):
            call_command("export_course_report", "XXX 999", stdout=StringIO())
