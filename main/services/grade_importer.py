import math
import re
import unicodedata

import openpyxl
from django.db import transaction

from main.models import Assessment, Student, StudentGrade


class GradeImporterService:
    STUDENT_NUMBER_HEADERS = ("student number", "student id", "ogrenci no", "id")
    GRADE_HEADERS = ("grade", "score", "not", "puan")

    @staticmethod
    def _normalize_header(value):
        normalized = unicodedata.normalize("NFKD", str(value))
        normalized = normalized.encode("ascii", "ignore").decode("ascii")
        return re.sub(r"[^a-z0-9]+", " ", normalized.casefold()).strip()

    @classmethod
    def _find_header_columns(cls, sheet):
        for row_number, row in enumerate(
            sheet.iter_rows(values_only=True), start=1
        ):
            student_number_column = None
            grade_column = None
            for column_number, value in enumerate(row):
                if value is None:
                    continue
                label = cls._normalize_header(value)
                if any(
                    re.search(rf"\b{re.escape(header)}\b", label)
                    for header in cls.STUDENT_NUMBER_HEADERS
                ):
                    student_number_column = column_number
                if any(
                    re.search(rf"\b{re.escape(header)}\b", label)
                    for header in cls.GRADE_HEADERS
                ):
                    grade_column = column_number
            if student_number_column is not None and grade_column is not None:
                return row_number, student_number_column, grade_column
        raise ValueError(
            "No header row containing student number and grade columns was found."
        )

    @staticmethod
    def _cell(row, column_number):
        if column_number >= len(row):
            return None
        return row[column_number]

    @staticmethod
    def import_excel_grades(assessment: Assessment, file_obj):
        """Import grade rows below a detected student-number/grade header."""
        workbook = openpyxl.load_workbook(file_obj, data_only=True)
        sheet = workbook.active
        header_row, student_number_column, grade_column = (
            GradeImporterService._find_header_columns(sheet)
        )

        rows = list(sheet.iter_rows(min_row=header_row + 1, values_only=True))
        if not rows:
            raise ValueError("The uploaded Excel file contains no grade rows.")

        student_numbers = []
        for row in rows:
            raw_student_number = GradeImporterService._cell(
                row, student_number_column
            )
            if raw_student_number is not None and str(raw_student_number).strip():
                student_numbers.append(str(raw_student_number).strip())

        students_map = {
            student.student_number: student
            for student in Student.objects.filter(
                student_number__in=student_numbers
            )
        }
        existing_grades = {
            grade.student_id: grade
            for grade in StudentGrade.objects.filter(
                assessment=assessment,
                student__student_number__in=student_numbers,
            )
        }

        grades_to_create = []
        grades_to_update = []
        errors = []

        for row_number, row in enumerate(rows, start=header_row + 1):
            raw_student_number = GradeImporterService._cell(
                row, student_number_column
            )
            raw_score = GradeImporterService._cell(row, grade_column)

            if raw_student_number is None and raw_score is None:
                continue
            student_number = (
                str(raw_student_number).strip()
                if raw_student_number is not None
                else ""
            )
            if not student_number:
                errors.append(f"Row {row_number}: Missing student number.")
                student = None
            else:
                student = students_map.get(student_number)
                if student is None:
                    errors.append(
                        f"Row {row_number}: Student with ID "
                        f"'{student_number}' does not exist in the database."
                    )

            try:
                score = float(raw_score)
            except (ValueError, TypeError):
                errors.append(
                    f"Row {row_number}: Invalid score value '{raw_score}'. "
                    "Must be a number."
                )
                continue

            if not math.isfinite(score) or score < 0.0 or score > 100.0:
                errors.append(
                    f"Row {row_number}: Score {score} must be between 0 and 100."
                )
                continue
            if student is None:
                continue

            if student.id in existing_grades:
                grade_record = existing_grades[student.id]
                grade_record.score = score
                grades_to_update.append(grade_record)
            else:
                grades_to_create.append(
                    StudentGrade(
                        student=student,
                        assessment=assessment,
                        score=score,
                    )
                )

        if errors:
            return {"success": False, "errors": errors, "created": 0, "updated": 0}

        with transaction.atomic():
            if grades_to_create:
                StudentGrade.objects.bulk_create(grades_to_create)
            if grades_to_update:
                StudentGrade.objects.bulk_update(grades_to_update, ["score"])

        return {
            "success": True,
            "errors": [],
            "created": len(grades_to_create),
            "updated": len(grades_to_update),
        }
