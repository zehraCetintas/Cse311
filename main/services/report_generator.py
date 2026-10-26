from io import BytesIO

import openpyxl
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from main.models import Course, ProgramOutcome, Student
from main.services.calculator import POCalculator


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
BELOW_FILL = PatternFill("solid", fgColor="F8CBAD")
ABOVE_FILL = PatternFill("solid", fgColor="C6EFCE")
TITLE_FONT = Font(bold=True, size=14)


class CourseReportService:
    DEFAULT_THRESHOLD = 60.0

    @staticmethod
    def build_course_report(course: Course, threshold: float = DEFAULT_THRESHOLD):
        """
        Collects everything an instructor needs to evaluate a course:
        each student's grades, LO and PO scores, class-level statistics,
        and the students who fall below the success threshold on any PO.
        """
        assessments = list(course.assessments.order_by("id"))
        lo_codes = list(course.learning_outcomes.order_by("code").values_list("code", flat=True))
        program_outcomes = list(
            ProgramOutcome.objects.filter(lo_mappings__learning_outcome__course=course)
            .distinct()
            .order_by("id")
        )
        po_codes = [po.code for po in program_outcomes]

        students = (
            Student.objects.filter(grades__assessment__course=course)
            .distinct()
            .order_by("student_number")
        )

        student_rows = []
        for student in students:
            grades_by_assessment = {
                grade.assessment_id: grade.score
                for grade in student.grades.filter(assessment__course=course)
            }
            lo_scores = POCalculator.calculate_student_lo_scores(student, course)
            po_scores = POCalculator.calculate_student_po_scores_for_course(student, course)
            student_rows.append(
                {
                    "student": student,
                    "grades": [grades_by_assessment.get(a.id) for a in assessments],
                    "lo_scores": lo_scores,
                    "po_scores": po_scores,
                }
            )

        lo_summary = [
            CourseReportService._summarize(
                code, None, [row["lo_scores"].get(code) for row in student_rows], threshold
            )
            for code in lo_codes
        ]
        po_summary = [
            CourseReportService._summarize(
                po.code, po.title, [row["po_scores"].get(po.code) for row in student_rows], threshold
            )
            for po in program_outcomes
        ]

        at_risk = []
        for row in student_rows:
            weak = {
                code: score
                for code, score in row["po_scores"].items()
                if score is not None and score < threshold
            }
            if weak:
                at_risk.append(
                    {
                        "student": row["student"],
                        "weak_pos": list(weak),
                        "lowest_score": min(weak.values()),
                    }
                )
        at_risk.sort(key=lambda item: item["lowest_score"])

        return {
            "course": course,
            "threshold": threshold,
            "assessments": assessments,
            "lo_codes": lo_codes,
            "po_codes": po_codes,
            "students": student_rows,
            "lo_summary": lo_summary,
            "po_summary": po_summary,
            "at_risk": at_risk,
        }

    @staticmethod
    def _summarize(code, title, scores, threshold):
        scores = [score for score in scores if score is not None]
        achieved = sum(1 for score in scores if score >= threshold)
        return {
            "code": code,
            "title": title,
            "count": len(scores),
            "average": round(sum(scores) / len(scores), 2) if scores else None,
            "min": min(scores) if scores else None,
            "max": max(scores) if scores else None,
            "achieved": achieved,
            "achievement_rate": round(100 * achieved / len(scores), 1) if scores else None,
        }

    @staticmethod
    def export_excel(course: Course, threshold: float = DEFAULT_THRESHOLD):
        """
        Builds the course report as an Excel (.xlsx) workbook and returns it
        as an in-memory file, ready to be saved or sent as a download.
        Sheets: Summary, Students, At-Risk Students.
        """
        report = CourseReportService.build_course_report(course, threshold)
        workbook = openpyxl.Workbook()

        CourseReportService._write_summary_sheet(workbook.active, report)
        CourseReportService._write_students_sheet(workbook.create_sheet("Students"), report)
        CourseReportService._write_at_risk_sheet(workbook.create_sheet("At-Risk Students"), report)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    @staticmethod
    def _write_summary_sheet(sheet, report):
        course = report["course"]
        threshold = report["threshold"]
        sheet.title = "Summary"

        sheet["A1"] = f"{course.code}: {course.title} - Course Outcome Report"
        sheet["A1"].font = TITLE_FONT
        sheet["A2"] = f"Instructor: {course.instructor.get_full_name() or course.instructor.username}" if course.instructor else "Instructor: -"
        sheet["A3"] = f"Generated: {timezone.localtime():%Y-%m-%d %H:%M}"
        sheet["A4"] = f"Students: {len(report['students'])}    Success threshold: {threshold:g}"

        headers = ["Outcome", "Title", "Average", "Min", "Max", f"Students >= {threshold:g}", "Achievement Rate (%)"]
        row = 6
        sheet.cell(row=row, column=1, value="Program Outcomes").font = Font(bold=True)
        row = CourseReportService._write_summary_table(sheet, row + 1, headers, report["po_summary"], threshold)

        sheet.cell(row=row + 1, column=1, value="Learning Outcomes").font = Font(bold=True)
        CourseReportService._write_summary_table(sheet, row + 2, headers, report["lo_summary"], threshold)

        CourseReportService._autosize(sheet, min_row=6)

    @staticmethod
    def _write_summary_table(sheet, start_row, headers, summary, threshold):
        CourseReportService._write_header(sheet, start_row, headers)
        row = start_row + 1
        for item in summary:
            values = [
                item["code"],
                item["title"] or "",
                item["average"],
                item["min"],
                item["max"],
                f"{item['achieved']}/{item['count']}",
                item["achievement_rate"],
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=value)
            CourseReportService._color_score(sheet.cell(row=row, column=3), threshold)
            row += 1
        return row

    @staticmethod
    def _write_students_sheet(sheet, report):
        threshold = report["threshold"]
        assessment_headers = [a.name for a in report["assessments"]]
        headers = (
            ["Student No", "Name"]
            + assessment_headers
            + report["lo_codes"]
            + report["po_codes"]
        )
        CourseReportService._write_header(sheet, 1, headers)
        score_start = 3 + len(assessment_headers)

        row = 2
        for student_row in report["students"]:
            student = student_row["student"]
            values = (
                [student.student_number, f"{student.first_name} {student.last_name}"]
                + student_row["grades"]
                + [student_row["lo_scores"].get(code) for code in report["lo_codes"]]
                + [student_row["po_scores"].get(code) for code in report["po_codes"]]
            )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                if column >= score_start:
                    CourseReportService._color_score(cell, threshold)
            row += 1

        # Class average row at the bottom
        if report["students"]:
            sheet.cell(row=row, column=1, value="Class Average").font = Font(bold=True)
            averages = (
                [None] * len(assessment_headers)
                + [item["average"] for item in report["lo_summary"]]
                + [item["average"] for item in report["po_summary"]]
            )
            for offset, value in enumerate(averages):
                cell = sheet.cell(row=row, column=3 + offset, value=value)
                cell.font = Font(bold=True)

        sheet.freeze_panes = "C2"
        CourseReportService._autosize(sheet)

    @staticmethod
    def _write_at_risk_sheet(sheet, report):
        headers = ["Student No", "Name", "POs Below Threshold", "Lowest PO Score"]
        CourseReportService._write_header(sheet, 1, headers)

        if not report["at_risk"]:
            sheet.cell(row=2, column=1, value="No students are below the threshold.")
        for row, item in enumerate(report["at_risk"], start=2):
            student = item["student"]
            sheet.cell(row=row, column=1, value=student.student_number)
            sheet.cell(row=row, column=2, value=f"{student.first_name} {student.last_name}")
            sheet.cell(row=row, column=3, value=", ".join(item["weak_pos"]))
            sheet.cell(row=row, column=4, value=item["lowest_score"]).fill = BELOW_FILL

        CourseReportService._autosize(sheet)

    @staticmethod
    def _write_header(sheet, row, headers):
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row, column=column, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

    @staticmethod
    def _color_score(cell, threshold):
        if isinstance(cell.value, (int, float)):
            cell.fill = ABOVE_FILL if cell.value >= threshold else BELOW_FILL

    @staticmethod
    def _autosize(sheet, min_row=1):
        for column_cells in sheet.iter_cols(min_row=min_row):
            length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(length + 2, 10), 50)
